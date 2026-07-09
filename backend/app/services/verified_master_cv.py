from __future__ import annotations

from functools import lru_cache
from importlib.resources import as_file, files
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from pathlib import Path
import re

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import Settings
from app.models.candidate_profile import CandidateProfile
from app.models.certification import Certification
from app.models.cv_version import CVVersion
from app.models.cv_version_artifact import CVVersionArtifact
from app.models.education import Education
from app.models.experience import Experience
from app.models.project import Project
from app.models.skill import Skill
from app.schemas.cv_version import CVMasterGenerationRequest
from app.services.cv_version import (
    CVVersionServiceError,
    DEFAULT_CANDIDATE_PROFILE_ID,
    clear_primary_for_type,
    current_utc_time,
    require_candidate_profile,
    validate_required_generated_artifacts,
    validate_parent_for_create,
)
from app.services.generated_cv_storage import (
    StoredGeneratedArtifact,
    remove_generated_artifact,
    resolve_generated_artifact_path,
    safe_download_filename,
    write_generated_artifact,
)


GENERATION_VERSION = "verified-master-cv-v1"
DOCX_RENDERER_VERSION = "python-docx-v1"
PDF_RENDERER_VERSION = "reportlab-packaged-noto-v2"
PDF_LATIN_FONT_NAME = "CareerPilotNotoSans"
PDF_CJK_FONT_NAME = "CareerPilotNotoSansSC"
PDF_LATIN_FONT_PACKAGE = "justmytype_western_core"
PDF_LATIN_FONT_RESOURCE = (
    "fonts/ofl/notosans/NotoSans[wdth,wght].ttf"
)
PDF_CJK_FONT_PACKAGE = "justmytype_intl_cjk"
PDF_CJK_FONT_RESOURCE = (
    "fonts/ofl/notosanssc/NotoSansSC[wght].ttf"
)
PDF_UNSUPPORTED_TEXT_MESSAGE = (
    "Generated CV contains characters unsupported by the PDF renderer."
)
PDF_LAYOUT_ERROR_MESSAGE = "Generated CV PDF layout values are invalid."

MONTH_ABBR = (
    "Jan",
    "Feb",
    "Mar",
    "Apr",
    "May",
    "Jun",
    "Jul",
    "Aug",
    "Sep",
    "Oct",
    "Nov",
    "Dec",
)

ARTIFACT_ORDER = {
    "docx": 0,
    "pdf": 1,
}


@dataclass(frozen=True)
class MasterCVLine:
    text: str
    style: str = "paragraph"


@dataclass(frozen=True)
class MasterCVDocument:
    lines: list[MasterCVLine]
    source_summary: dict[str, int]

    @property
    def plain_lines(self) -> list[str]:
        return [line.text for line in self.lines if line.text]


@dataclass(frozen=True)
class GeneratedMasterCVResult:
    version: CVVersion
    artifacts: list[CVVersionArtifact] = field(default_factory=list)
    created: bool = False


def normalize_whitespace(value: str | None) -> str | None:
    if value is None:
        return None

    lines = []
    for line in value.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        cleaned = re.sub(r"[ \t\f\v]+", " ", line).strip()
        if cleaned:
            lines.append(cleaned)

    return "\n".join(lines) or None


def split_clean_lines(value: str | None) -> list[str]:
    normalized = normalize_whitespace(value)
    return normalized.split("\n") if normalized else []


def clean_list(values: list[str] | None) -> list[str]:
    cleaned: list[str] = []
    seen: set[str] = set()

    for value in values or []:
        normalized = normalize_whitespace(value)
        if normalized is None:
            continue

        key = normalized.casefold()
        if key in seen:
            continue

        seen.add(key)
        cleaned.append(normalized)

    return cleaned


def format_month_year(value: date | None) -> str | None:
    if value is None:
        return None

    return f"{MONTH_ABBR[value.month - 1]} {value.year}"


def format_date_range(
    start_date: date | None,
    end_date: date | None,
    *,
    is_current: bool = False,
) -> str | None:
    start = format_month_year(start_date)

    if is_current:
        return f"{start} - Present" if start else "Present"

    end = format_month_year(end_date)

    if start and end:
        return f"{start} - {end}"

    return start or end


def join_parts(parts: list[str | None], separator: str = " | ") -> str | None:
    cleaned = [part for part in parts if part]
    return separator.join(cleaned) if cleaned else None


def reverse_chronological_key(
    start_date: date | None,
    end_date: date | None,
    is_current: bool,
    record_id: int,
) -> tuple[date, date, int]:
    comparable_end = date.max if is_current else end_date or start_date
    comparable_start = start_date or date.min

    return (
        comparable_end or date.min,
        comparable_start,
        -record_id,
    )


def add_line(
    lines: list[MasterCVLine],
    text: str | None,
    *,
    style: str = "paragraph",
) -> None:
    normalized = normalize_whitespace(text)

    if normalized:
        for line in normalized.split("\n"):
            lines.append(MasterCVLine(line, style))


def add_section_heading(
    lines: list[MasterCVLine],
    heading: str,
) -> None:
    if lines:
        lines.append(MasterCVLine("", "spacer"))

    lines.append(MasterCVLine(heading, "section_heading"))


def add_contact_section(
    lines: list[MasterCVLine],
    profile: CandidateProfile,
) -> None:
    """Header block: name, title, and a single inline contact line.

    Contact details are joined with a middot separator for a modern
    look. Only fields the candidate actually provided are shown —
    nothing is invented.
    """
    add_line(lines, profile.full_name, style="name")
    add_line(lines, profile.professional_title, style="subtitle")

    contact_values = [
        normalize_whitespace(value)
        for value in (
            profile.email,
            profile.phone,
            profile.location,
            profile.linkedin_url,
            profile.github_url,
            profile.portfolio_url,
        )
    ]
    present = [value for value in contact_values if value]
    if present:
        lines.append(
            MasterCVLine("  ·  ".join(present), "contact")
        )


def add_summary_section(
    lines: list[MasterCVLine],
    profile: CandidateProfile,
) -> bool:
    summary_lines = split_clean_lines(profile.professional_summary)
    if not summary_lines:
        return False

    add_section_heading(lines, "Professional Summary")
    for summary_line in summary_lines:
        add_line(lines, summary_line)

    return True


def add_skills_section(
    lines: list[MasterCVLine],
    skills: list[Skill],
) -> bool:
    grouped: dict[str, list[str]] = {}
    seen_names: set[str] = set()

    for skill in sorted(
        skills,
        key=lambda item: (
            normalize_whitespace(item.category or "") or "",
            normalize_whitespace(item.name) or "",
            item.id,
        ),
    ):
        name = normalize_whitespace(skill.name)
        if name is None:
            continue

        key = name.casefold()
        if key in seen_names:
            continue

        seen_names.add(key)
        category = normalize_whitespace(skill.category) or "Skills"
        proficiency = normalize_whitespace(skill.proficiency)
        display_name = (
            f"{name} ({proficiency})" if proficiency else name
        )
        grouped.setdefault(category, []).append(display_name)

    if not grouped:
        return False

    add_section_heading(lines, "Skills")
    for category in sorted(grouped):
        lines.append(
            MasterCVLine(
                f"{category}: {', '.join(grouped[category])}",
                "paragraph",
            )
        )

    return True


def add_experience_section(
    lines: list[MasterCVLine],
    experiences: list[Experience],
) -> bool:
    ordered = sorted(
        experiences,
        key=lambda item: reverse_chronological_key(
            item.start_date,
            item.end_date,
            item.is_current,
            item.id,
        ),
        reverse=True,
    )
    if not ordered:
        return False

    add_section_heading(lines, "Work Experience")
    for experience in ordered:
        title = join_parts(
            [experience.job_title, experience.company],
            separator=", ",
        )
        add_line(lines, title, style="entry_heading")
        add_line(
            lines,
            join_parts(
                [
                    experience.location,
                    experience.employment_type,
                    format_date_range(
                        experience.start_date,
                        experience.end_date,
                        is_current=experience.is_current,
                    ),
                ]
            ),
            style="meta",
        )

        for description_line in split_clean_lines(experience.description):
            add_line(lines, description_line)

        for value in clean_list(experience.responsibilities):
            add_line(lines, value, style="bullet")

        for value in clean_list(experience.achievements):
            add_line(lines, value, style="bullet")

        technologies = clean_list(experience.technologies)
        if technologies:
            add_line(
                lines,
                f"Technologies: {', '.join(technologies)}",
                style="meta",
            )

    return True


def add_education_section(
    lines: list[MasterCVLine],
    education_records: list[Education],
) -> bool:
    ordered = sorted(
        education_records,
        key=lambda item: reverse_chronological_key(
            item.start_date,
            item.end_date,
            item.is_current,
            item.id,
        ),
        reverse=True,
    )
    if not ordered:
        return False

    add_section_heading(lines, "Education")
    for education in ordered:
        title = join_parts(
            [
                join_parts(
                    [education.degree, education.field_of_study],
                    separator=" in ",
                ),
                education.institution,
            ],
            separator=", ",
        )
        add_line(lines, title, style="entry_heading")
        add_line(
            lines,
            join_parts(
                [
                    education.location,
                    format_date_range(
                        education.start_date,
                        education.end_date,
                        is_current=education.is_current,
                    ),
                ]
            ),
            style="meta",
        )
        add_line(lines, education.grade, style="meta")

        for description_line in split_clean_lines(education.description):
            add_line(lines, description_line)

    return True


def add_projects_section(
    lines: list[MasterCVLine],
    projects: list[Project],
) -> bool:
    ordered = sorted(
        projects,
        key=lambda item: reverse_chronological_key(
            item.start_date,
            item.end_date,
            item.is_current,
            item.id,
        ),
        reverse=True,
    )
    if not ordered:
        return False

    add_section_heading(lines, "Projects")
    for project in ordered:
        title = join_parts([project.name, project.role], separator=" - ")
        add_line(lines, title, style="entry_heading")
        add_line(
            lines,
            join_parts(
                [
                    project.project_type,
                    format_date_range(
                        project.start_date,
                        project.end_date,
                        is_current=project.is_current,
                    ),
                ]
            ),
            style="meta",
        )

        for description_line in split_clean_lines(project.description):
            add_line(lines, description_line)

        technologies = clean_list(project.technologies)
        if technologies:
            add_line(
                lines,
                f"Technologies: {', '.join(technologies)}",
                style="meta",
            )

        for value in clean_list(project.features):
            add_line(lines, value, style="bullet")

        for value in clean_list(project.achievements):
            add_line(lines, value, style="bullet")

        add_line(lines, project.project_url, style="meta")
        add_line(lines, project.github_url, style="meta")

    return True


def add_certifications_section(
    lines: list[MasterCVLine],
    certifications: list[Certification],
) -> bool:
    ordered = sorted(
        certifications,
        key=lambda item: (
            item.issue_date or date.min,
            -item.id,
        ),
        reverse=True,
    )
    if not ordered:
        return False

    add_section_heading(lines, "Certifications")
    for certification in ordered:
        add_line(
            lines,
            join_parts(
                [
                    certification.name,
                    certification.issuing_organization,
                ],
                separator=", ",
            ),
            style="entry_heading",
        )
        expiry = (
            "No expiry"
            if certification.does_not_expire
            else f"Expires {format_month_year(certification.expiry_date)}"
            if certification.expiry_date
            else None
        )
        add_line(
            lines,
            join_parts(
                [
                    f"Issued {format_month_year(certification.issue_date)}",
                    expiry,
                ]
            ),
            style="meta",
        )
        credential_id = normalize_whitespace(certification.credential_id)
        if credential_id:
            add_line(lines, f"Credential ID: {credential_id}", style="meta")
        add_line(lines, certification.credential_url, style="meta")

        for description_line in split_clean_lines(certification.description):
            add_line(lines, description_line)

        certification_skills = clean_list(certification.skills)
        if certification_skills:
            add_line(
                lines,
                f"Skills: {', '.join(certification_skills)}",
                style="meta",
            )

    return True


def select_verified_records(
    database: Session,
    candidate_profile_id: int,
) -> tuple[
    list[Education],
    list[Experience],
    list[Project],
    list[Skill],
    list[Certification],
]:
    education_records = list(
        database.scalars(
            select(Education).where(
                Education.candidate_profile_id == candidate_profile_id,
                Education.verified.is_(True),
            )
        ).all()
    )
    experiences = list(
        database.scalars(
            select(Experience).where(
                Experience.candidate_profile_id == candidate_profile_id,
                Experience.verified.is_(True),
            )
        ).all()
    )
    projects = list(
        database.scalars(
            select(Project).where(
                Project.candidate_profile_id == candidate_profile_id,
                Project.verified.is_(True),
            )
        ).all()
    )
    skills = list(
        database.scalars(
            select(Skill).where(
                Skill.candidate_profile_id == candidate_profile_id,
                Skill.verified.is_(True),
            )
        ).all()
    )
    certifications = list(
        database.scalars(
            select(Certification).where(
                Certification.candidate_profile_id == candidate_profile_id,
                Certification.verified.is_(True),
            )
        ).all()
    )

    return (
        education_records,
        experiences,
        projects,
        skills,
        certifications,
    )


def build_master_cv_document(
    profile: CandidateProfile,
    *,
    education_records: list[Education],
    experiences: list[Experience],
    projects: list[Project],
    skills: list[Skill],
    certifications: list[Certification],
) -> MasterCVDocument:
    lines: list[MasterCVLine] = []
    summary: dict[str, int] = {
        "profile": 1,
        "education": len(education_records),
        "work_experiences": len(experiences),
        "projects": len(projects),
        "skills": len(skills),
        "certifications": len(certifications),
    }
    populated_sections = 0

    add_contact_section(lines, profile)

    if add_summary_section(lines, profile):
        populated_sections += 1

    for section_added in (
        add_experience_section(lines, experiences),
        add_education_section(lines, education_records),
        add_skills_section(lines, skills),
        add_projects_section(lines, projects),
        add_certifications_section(lines, certifications),
    ):
        if section_added:
            populated_sections += 1

    if populated_sections == 0:
        raise CVVersionServiceError(
            "No eligible verified facts are available for a master CV.",
            status_code=400,
        )

    return MasterCVDocument(lines=lines, source_summary=summary)


def require_verified_candidate_profile(
    database: Session,
    candidate_profile_id: int,
) -> CandidateProfile:
    profile = require_candidate_profile(database, candidate_profile_id)

    if not profile.facts_verified:
        raise CVVersionServiceError(
            "Candidate profile facts must be verified before generation.",
            status_code=400,
        )

    return profile


def validate_generation_parent(parent: CVVersion) -> None:
    if parent.status != "ready":
        raise CVVersionServiceError(
            "Only ready parsed or verified master CV versions can be used.",
            status_code=400,
        )


# Professional accent palette used across both renderers. Navy for
# the name and section headings, slate grey for supporting text.
CV_ACCENT_RGB = (0x1F, 0x3A, 0x5F)
CV_MUTED_RGB = (0x55, 0x5F, 0x6B)


def _set_docx_heading_rule(paragraph: object) -> None:
    """Add a thin bottom border under a paragraph (section rule)."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "1F3A5F")
    borders.append(bottom)
    p_pr.append(borders)


def render_docx(
    document_data: MasterCVDocument,
    output_path: Path,
) -> int | None:
    from docx import Document as DocxDocument
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.shared import Inches, Pt, RGBColor

    accent = RGBColor(*CV_ACCENT_RGB)
    muted = RGBColor(*CV_MUTED_RGB)

    document = DocxDocument()

    for section in document.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10.5)

    for line in document_data.lines:
        if line.style == "spacer":
            continue

        if line.style == "section_heading":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(5)
            run = paragraph.add_run(line.text)
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = accent
            _set_docx_heading_rule(paragraph)
            continue

        if line.style == "bullet":
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(line.text)
            paragraph.paragraph_format.space_after = Pt(2)
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)

        if line.style == "name":
            paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            paragraph.paragraph_format.space_after = Pt(1)
            run = paragraph.add_run(line.text)
            run.bold = True
            run.font.size = Pt(22)
            run.font.color.rgb = accent
        elif line.style == "subtitle":
            run = paragraph.add_run(line.text)
            run.italic = True
            run.font.size = Pt(12)
            run.font.color.rgb = muted
        elif line.style == "contact":
            run = paragraph.add_run(line.text)
            run.font.size = Pt(9.5)
            run.font.color.rgb = muted
        elif line.style == "entry_heading":
            run = paragraph.add_run(line.text)
            run.bold = True
            run.font.size = Pt(11)
        elif line.style == "meta":
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(line.text)
            run.italic = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = muted
        else:
            paragraph.add_run(line.text)

    document.save(output_path)
    return None


@dataclass(frozen=True)
class PDFTextRun:
    font_name: str
    text: str


def packaged_font_path(package_name: str, resource_path: str) -> Path:
    resource = files(package_name).joinpath(resource_path)
    with as_file(resource) as font_path:
        return Path(font_path)


@lru_cache(maxsize=2)
def font_codepoints(package_name: str, resource_path: str) -> frozenset[int]:
    from fontTools.ttLib import TTFont

    font_path = packaged_font_path(package_name, resource_path)
    font = TTFont(font_path)
    codepoints: set[int] = set()

    for table in font["cmap"].tables:
        if table.isUnicode():
            codepoints.update(table.cmap)

    return frozenset(codepoints)


def latin_font_codepoints() -> frozenset[int]:
    return font_codepoints(
        PDF_LATIN_FONT_PACKAGE,
        PDF_LATIN_FONT_RESOURCE,
    )


def cjk_font_codepoints() -> frozenset[int]:
    return font_codepoints(
        PDF_CJK_FONT_PACKAGE,
        PDF_CJK_FONT_RESOURCE,
    )


def pdf_font_for_character(value: str) -> str:
    if value.isspace() or ord(value) in latin_font_codepoints():
        return PDF_LATIN_FONT_NAME

    if ord(value) in cjk_font_codepoints():
        return PDF_CJK_FONT_NAME

    raise CVVersionServiceError(
        PDF_UNSUPPORTED_TEXT_MESSAGE,
        status_code=400,
    )


def preflight_pdf_text(document_data: MasterCVDocument) -> None:
    for line in document_data.lines:
        for character in line.text:
            pdf_font_for_character(character)


def register_pdf_fonts() -> None:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if PDF_LATIN_FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(
            TTFont(
                PDF_LATIN_FONT_NAME,
                str(
                    packaged_font_path(
                        PDF_LATIN_FONT_PACKAGE,
                        PDF_LATIN_FONT_RESOURCE,
                    )
                ),
            )
        )

    if PDF_CJK_FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(
            TTFont(
                PDF_CJK_FONT_NAME,
                str(
                    packaged_font_path(
                        PDF_CJK_FONT_PACKAGE,
                        PDF_CJK_FONT_RESOURCE,
                    )
                ),
            )
        )


def split_pdf_text_runs(text: str) -> list[PDFTextRun]:
    runs: list[PDFTextRun] = []
    current_font: str | None = None
    current_text = ""

    for character in text:
        font_name = pdf_font_for_character(character)

        if current_text and font_name != current_font:
            assert current_font is not None
            runs.append(PDFTextRun(current_font, current_text))
            current_text = ""

        current_font = font_name
        current_text += character

    if current_text:
        assert current_font is not None
        runs.append(PDFTextRun(current_font, current_text))

    return runs


def pdf_text_width(
    runs: list[PDFTextRun],
    font_size: float,
) -> float:
    from reportlab.pdfbase.pdfmetrics import stringWidth

    return sum(
        stringWidth(run.text, run.font_name, font_size)
        for run in runs
    )


CV_ACCENT_PDF = (
    CV_ACCENT_RGB[0] / 255,
    CV_ACCENT_RGB[1] / 255,
    CV_ACCENT_RGB[2] / 255,
)
CV_MUTED_PDF = (
    CV_MUTED_RGB[0] / 255,
    CV_MUTED_RGB[1] / 255,
    CV_MUTED_RGB[2] / 255,
)
CV_BLACK_PDF = (0.0, 0.0, 0.0)


def render_pdf_text_runs(
    pdf: object,
    runs: list[PDFTextRun],
    *,
    x: float,
    y: float,
    font_size: float,
    color: tuple[float, float, float] = CV_BLACK_PDF,
) -> None:
    text = pdf.beginText(x, y)
    text.setFillColorRGB(*color)

    for run in runs:
        text.setFont(run.font_name, font_size)
        text.textOut(run.text)

    pdf.drawText(text)


def wrap_pdf_text(
    text: str,
    *,
    font_size: float,
    max_width: float,
) -> list[list[PDFTextRun]]:
    def split_word(word: str) -> list[list[PDFTextRun]]:
        pieces: list[list[PDFTextRun]] = []
        current_piece = ""

        for character in word:
            candidate = f"{current_piece}{character}"
            candidate_runs = split_pdf_text_runs(candidate)
            if (
                not current_piece
                or pdf_text_width(candidate_runs, font_size) <= max_width
            ):
                current_piece = candidate
                continue

            pieces.append(split_pdf_text_runs(current_piece))
            current_piece = character

        if current_piece:
            pieces.append(split_pdf_text_runs(current_piece))

        return pieces

    words = text.split()
    if not words:
        return [[]]

    wrapped: list[list[PDFTextRun]] = []
    current = ""

    for word in words:
        candidates = split_word(word)

        for candidate_runs in candidates:
            candidate_word = "".join(run.text for run in candidate_runs)
            candidate = (
                f"{current} {candidate_word}"
                if current
                else candidate_word
            )
            candidate_line_runs = split_pdf_text_runs(candidate)
            if (
                not current
                or pdf_text_width(candidate_line_runs, font_size) <= max_width
            ):
                current = candidate
                continue

            wrapped.append(split_pdf_text_runs(current))
            current = candidate_word

    if current:
        wrapped.append(split_pdf_text_runs(current))

    return wrapped


def render_pdf(
    document_data: MasterCVDocument,
    output_path: Path,
) -> int | None:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.units import inch
    from reportlab.pdfgen import canvas

    preflight_pdf_text(document_data)
    register_pdf_fonts()

    pdf = canvas.Canvas(
        str(output_path),
        pagesize=LETTER,
        initialFontName=PDF_LATIN_FONT_NAME,
    )
    width, height = LETTER
    margin = 0.75 * inch
    y = height - margin
    page_count = 1
    printable_height = height - (2 * margin)

    def start_new_page() -> None:
        nonlocal y, page_count

        pdf.showPage()
        page_count += 1
        y = height - margin

    def ensure_space(required_height: float) -> None:
        if required_height <= 0 or required_height > printable_height:
            raise CVVersionServiceError(
                PDF_LAYOUT_ERROR_MESSAGE,
                status_code=500,
            )

        if y - required_height < margin:
            start_new_page()

        if y - required_height < margin:
            raise CVVersionServiceError(
                PDF_LAYOUT_ERROR_MESSAGE,
                status_code=500,
            )

    def apply_spacing(spacing: float) -> None:
        nonlocal y

        if spacing < 0:
            raise CVVersionServiceError(
                PDF_LAYOUT_ERROR_MESSAGE,
                status_code=500,
            )

        if spacing == 0:
            return

        if y - spacing >= margin:
            y -= spacing
        else:
            y = margin

    def draw_wrapped(
        text: str,
        *,
        font_size: float,
        line_height: float,
        left: float,
        space_after: float,
        keep_with_next_height: float = 0,
        color: tuple[float, float, float] = CV_BLACK_PDF,
        rule: bool = False,
    ) -> None:
        nonlocal y

        if font_size <= 0 or line_height <= 0:
            raise CVVersionServiceError(
                PDF_LAYOUT_ERROR_MESSAGE,
                status_code=500,
            )

        max_width = width - margin - left
        if max_width <= 0:
            raise CVVersionServiceError(
                PDF_LAYOUT_ERROR_MESSAGE,
                status_code=500,
            )

        wrapped_lines = wrap_pdf_text(
            text,
            font_size=font_size,
            max_width=max_width,
        )
        keep_together_height = (
            line_height * len(wrapped_lines)
            + space_after
            + keep_with_next_height
        )
        if (
            keep_with_next_height > 0
            and keep_together_height <= printable_height
            and y - keep_together_height < margin
        ):
            start_new_page()

        for wrapped_line_runs in wrapped_lines:
            ensure_space(line_height)
            render_pdf_text_runs(
                pdf,
                wrapped_line_runs,
                x=left,
                y=y,
                font_size=font_size,
                color=color,
            )
            y -= line_height

        if rule:
            ensure_space(3)
            pdf.setStrokeColorRGB(*CV_ACCENT_PDF)
            pdf.setLineWidth(0.6)
            pdf.line(margin, y + 3, width - margin, y + 3)
            pdf.setStrokeColorRGB(*CV_BLACK_PDF)

        apply_spacing(space_after)

    for line in document_data.lines:
        if line.style == "spacer":
            apply_spacing(6)
            continue

        if line.style == "section_heading":
            draw_wrapped(
                line.text,
                font_size=12,
                line_height=15,
                left=margin,
                space_after=6,
                keep_with_next_height=12,
                color=CV_ACCENT_PDF,
                rule=True,
            )
        elif line.style == "name":
            draw_wrapped(
                line.text,
                font_size=20,
                line_height=23,
                left=margin,
                space_after=2,
                color=CV_ACCENT_PDF,
            )
        elif line.style == "subtitle":
            draw_wrapped(
                line.text,
                font_size=11.5,
                line_height=14,
                left=margin,
                space_after=2,
                color=CV_MUTED_PDF,
            )
        elif line.style == "contact":
            draw_wrapped(
                line.text,
                font_size=9.5,
                line_height=12,
                left=margin,
                space_after=1,
                color=CV_MUTED_PDF,
            )
        elif line.style == "entry_heading":
            draw_wrapped(
                line.text,
                font_size=11,
                line_height=13,
                left=margin,
                space_after=1,
                keep_with_next_height=12,
            )
        elif line.style == "meta":
            draw_wrapped(
                line.text,
                font_size=9.5,
                line_height=12,
                left=margin,
                space_after=3,
                color=CV_MUTED_PDF,
            )
        elif line.style == "bullet":
            draw_wrapped(
                line.text,
                font_size=10,
                line_height=12.5,
                left=margin + 14,
                space_after=2,
            )
        else:
            draw_wrapped(
                line.text,
                font_size=10,
                line_height=12.5,
                left=margin,
                space_after=2,
            )

    pdf.save()
    return page_count


def list_cv_version_artifacts(
    version_id: int,
    database: Session,
    *,
    candidate_profile_id: int = DEFAULT_CANDIDATE_PROFILE_ID,
) -> list[CVVersionArtifact]:
    return sorted(
        database.scalars(
            select(CVVersionArtifact)
            .join(CVVersion, CVVersion.id == CVVersionArtifact.cv_version_id)
            .where(
                CVVersionArtifact.cv_version_id == version_id,
                CVVersion.candidate_profile_id == candidate_profile_id,
            )
        ).all(),
        key=lambda artifact: (
            ARTIFACT_ORDER.get(artifact.artifact_type, 99),
            artifact.id,
        ),
    )


def get_cv_version_artifact_or_404(
    artifact_id: int,
    database: Session,
    *,
    candidate_profile_id: int = DEFAULT_CANDIDATE_PROFILE_ID,
) -> CVVersionArtifact:
    artifact = database.scalars(
        select(CVVersionArtifact)
        .join(CVVersion, CVVersion.id == CVVersionArtifact.cv_version_id)
        .where(
            CVVersionArtifact.id == artifact_id,
            CVVersion.candidate_profile_id == candidate_profile_id,
        )
    ).first()

    if artifact is None:
        raise CVVersionServiceError(
            "Generated CV artifact was not found.",
            status_code=404,
        )

    return artifact


def resolve_artifact_download_path(
    artifact: CVVersionArtifact,
) -> Path:
    try:
        return resolve_generated_artifact_path(artifact.storage_path)
    except Exception as exc:
        raise CVVersionServiceError(
            "Generated CV artifact file was not found.",
            status_code=404,
        ) from exc


def find_existing_generation(
    database: Session,
    *,
    candidate_profile_id: int,
    parent_version_id: int,
    generation_key: str,
) -> CVVersion | None:
    return database.scalars(
        select(CVVersion)
        .where(
            CVVersion.candidate_profile_id == candidate_profile_id,
            CVVersion.parent_version_id == parent_version_id,
            CVVersion.version_type == "verified_master",
            CVVersion.generation_key == generation_key,
        )
        .order_by(CVVersion.id.desc())
    ).first()


def create_draft_generation_version(
    database: Session,
    *,
    profile: CandidateProfile,
    parent: CVVersion,
    request: CVMasterGenerationRequest,
    source_summary: dict[str, int],
) -> CVVersion:
    now = current_utc_time()
    version = CVVersion(
        candidate_profile_id=profile.id,
        source_document_id=None,
        parent_version_id=parent.id,
        version_type="verified_master",
        status="draft",
        display_name=(
            request.display_name
            or f"{profile.full_name} Verified Master CV"
        ),
        description=request.description,
        job_reference=None,
        is_primary=False,
        storage_path=None,
        mime_type=None,
        file_size=None,
        sha256=None,
        generation_key=request.idempotency_key,
        generation_version=GENERATION_VERSION,
        source_summary=source_summary,
        generated_at=None,
        created_at=now,
        updated_at=now,
        archived_at=None,
    )

    try:
        database.add(version)
        database.commit()
        database.refresh(version)
    except Exception as exc:
        database.rollback()
        raise CVVersionServiceError(
            "Failed to create verified master CV version.",
            status_code=500,
        ) from exc

    return version


def mark_generation_failed(
    version_id: int,
    database: Session,
) -> None:
    try:
        database.rollback()
        version = database.get(CVVersion, version_id)
        if version is None:
            return

        version.status = "failed"
        version.is_primary = False
        version.updated_at = current_utc_time()
        database.commit()
    except Exception:
        database.rollback()


def cleanup_stored_artifacts(
    artifacts: list[StoredGeneratedArtifact],
) -> None:
    for artifact in artifacts:
        try:
            remove_generated_artifact(artifact.storage_path)
        except Exception:
            continue


def validate_stored_artifacts_for_ready(
    stored_artifacts: list[StoredGeneratedArtifact],
) -> None:
    artifact_types = [artifact.artifact_type for artifact in stored_artifacts]

    if sorted(artifact_types) != ["docx", "pdf"]:
        raise CVVersionServiceError(
            (
                "Verified master generation requires one DOCX artifact "
                "and one PDF artifact."
            ),
            status_code=500,
        )


def persist_generated_artifacts(
    version: CVVersion,
    stored_artifacts: list[StoredGeneratedArtifact],
    database: Session,
    *,
    request: CVMasterGenerationRequest,
    profile: CandidateProfile,
) -> list[CVVersionArtifact]:
    now = current_utc_time()
    persisted: list[CVVersionArtifact] = []

    try:
        validate_stored_artifacts_for_ready(stored_artifacts)

        version.is_primary = False

        for stored_artifact in stored_artifacts:
            renderer_version = (
                DOCX_RENDERER_VERSION
                if stored_artifact.artifact_type == "docx"
                else PDF_RENDERER_VERSION
            )
            artifact = CVVersionArtifact(
                cv_version_id=version.id,
                artifact_type=stored_artifact.artifact_type,
                mime_type=stored_artifact.mime_type,
                safe_filename=stored_artifact.safe_filename,
                storage_path=stored_artifact.storage_path,
                file_size=stored_artifact.file_size,
                sha256=stored_artifact.sha256,
                page_count=stored_artifact.page_count,
                renderer_version=renderer_version,
                content_version=GENERATION_VERSION,
                created_at=now,
            )
            database.add(artifact)
            persisted.append(artifact)

        database.flush()
        validate_required_generated_artifacts(version.id, database)

        version.status = "ready"
        version.generated_at = now
        version.updated_at = now

        if request.is_primary:
            clear_primary_for_type(
                database,
                candidate_profile_id=profile.id,
                version_type=version.version_type,
                now=now,
            )
            version.is_primary = True

        database.commit()
        database.refresh(version)
        for artifact in persisted:
            database.refresh(artifact)
    except CVVersionServiceError:
        database.rollback()
        cleanup_stored_artifacts(stored_artifacts)
        mark_generation_failed(version.id, database)
        raise
    except Exception as exc:
        database.rollback()
        cleanup_stored_artifacts(stored_artifacts)
        mark_generation_failed(version.id, database)
        raise CVVersionServiceError(
            "Failed to save verified master CV artifacts.",
            status_code=500,
        ) from exc

    return sorted(
        persisted,
        key=lambda artifact: (
            ARTIFACT_ORDER.get(artifact.artifact_type, 99),
            artifact.id,
        ),
    )


def generate_verified_master_cv(
    parent_version_id: int,
    request: CVMasterGenerationRequest,
    database: Session,
    *,
    candidate_profile_id: int = DEFAULT_CANDIDATE_PROFILE_ID,
    settings: Settings | None = None,
) -> GeneratedMasterCVResult:
    profile = require_verified_candidate_profile(
        database,
        candidate_profile_id,
    )
    parent = validate_parent_for_create(
        "verified_master",
        parent_version_id,
        database,
        candidate_profile_id=candidate_profile_id,
    )
    assert parent is not None
    validate_generation_parent(parent)

    existing = find_existing_generation(
        database,
        candidate_profile_id=candidate_profile_id,
        parent_version_id=parent.id,
        generation_key=request.idempotency_key,
    )
    if existing is not None:
        return GeneratedMasterCVResult(
            version=existing,
            artifacts=list_cv_version_artifacts(
                existing.id,
                database,
                candidate_profile_id=candidate_profile_id,
            ),
            created=False,
        )

    (
        education_records,
        experiences,
        projects,
        skills,
        certifications,
    ) = select_verified_records(database, candidate_profile_id)

    document_data = build_master_cv_document(
        profile,
        education_records=education_records,
        experiences=experiences,
        projects=projects,
        skills=skills,
        certifications=certifications,
    )
    version = create_draft_generation_version(
        database,
        profile=profile,
        parent=parent,
        request=request,
        source_summary=document_data.source_summary,
    )
    stored: list[StoredGeneratedArtifact] = []

    try:
        preflight_pdf_text(document_data)

        docx_safe_filename = safe_download_filename(
            profile.full_name,
            version.id,
            "docx",
        )
        stored.append(
            write_generated_artifact(
                "docx",
                docx_safe_filename,
                lambda path: render_docx(document_data, path),
                settings=settings,
            )
        )

        pdf_safe_filename = safe_download_filename(
            profile.full_name,
            version.id,
            "pdf",
        )
        stored.append(
            write_generated_artifact(
                "pdf",
                pdf_safe_filename,
                lambda path: render_pdf(document_data, path),
                settings=settings,
            )
        )
    except CVVersionServiceError:
        cleanup_stored_artifacts(stored)
        mark_generation_failed(version.id, database)
        raise
    except Exception as exc:
        cleanup_stored_artifacts(stored)
        mark_generation_failed(version.id, database)
        raise CVVersionServiceError(
            "Failed to generate verified master CV artifacts.",
            status_code=500,
        ) from exc

    artifacts = persist_generated_artifacts(
        version,
        stored,
        database,
        request=request,
        profile=profile,
    )

    return GeneratedMasterCVResult(
        version=version,
        artifacts=artifacts,
        created=True,
    )


def delete_generated_artifact_files(
    artifacts: list[CVVersionArtifact],
) -> None:
    for artifact in artifacts:
        remove_generated_artifact(artifact.storage_path)
