from datetime import date, datetime, timezone
from hashlib import sha256
from pathlib import Path
from uuid import uuid4

import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, inspect, select, text
from sqlalchemy.orm import Session


@pytest.fixture(autouse=True)
def clear_phase6_rows_after_test(db_session: Session) -> None:
    yield
    db_session.rollback()
    db_session.execute(text("PRAGMA foreign_keys=OFF"))
    for table_name in (
        "job_tailored_cvs",
        "cv_version_artifacts",
        "cv_versions",
        "job_skill_requirements",
        "normalized_jobs",
        "job_sources",
        "work_experiences",
        "projects",
        "skills",
        "candidate_profiles",
    ):
        db_session.execute(text(f"DELETE FROM {table_name}"))
    db_session.commit()
    db_session.execute(text("PRAGMA foreign_keys=ON"))

from app.models.candidate_profile import CandidateProfile
from app.models.cv_version import CVVersion
from app.models.experience import Experience
from app.models.job import JobSkillRequirement, NormalizedJob
from app.models.job_tailored_cv import JobTailoredCV
from app.models.project import Project
from app.models.skill import Skill
from app.schemas.cv_version import CVMasterGenerationRequest
from app.schemas.job import (
    JobSkillRequirementCreate,
    JobSourceCreate,
    NormalizedJobCreate,
)
from app.services.cv_tailoring import (
    CVTailoringError,
    extract_job_keywords,
    generate_tailored_cv,
    get_tailored_cv_or_404,
    reorder_skills,
    select_relevant_experiences,
)
from app.services.job import create_job_source, create_normalized_job


PHASE5_HEAD = "f2a3b4c5d6e7"
PHASE6_HEAD = "a3b4c5d6e7f8"


def content_hash(seed: str | None = None) -> str:
    return sha256((seed or uuid4().hex).encode()).hexdigest()


def make_profile(database: Session) -> CandidateProfile:
    profile = CandidateProfile(
        id=1,
        full_name="Pavan Test",
        email="pavan@example.com",
        location="London",
        professional_title="Senior Backend Engineer",
        professional_summary=(
            "Backend engineer focused on reliable APIs."
        ),
        work_authorization="UK citizen",
        requires_sponsorship=False,
        willing_to_relocate=False,
        preferred_work_modes=["remote"],
        target_roles=["Backend Engineer"],
        preferred_locations=["London"],
        salary_currency="GBP",
        minimum_match_score=50,
        auto_apply_enabled=False,
        facts_verified=True,
    )
    database.add(profile)
    database.commit()
    return profile


def add_skill(
    database: Session,
    name: str,
    *,
    verified: bool = True,
) -> Skill:
    skill = Skill(
        candidate_profile_id=1,
        name=name,
        proficiency="advanced",
        verified=verified,
    )
    database.add(skill)
    database.commit()
    return skill


def add_experience(
    database: Session,
    company: str,
    technologies: list[str],
    start: date,
    *,
    verified: bool = True,
) -> Experience:
    experience = Experience(
        candidate_profile_id=1,
        company=company,
        job_title="Engineer",
        start_date=start,
        end_date=None,
        is_current=True,
        responsibilities=["Verified responsibility."],
        achievements=[],
        technologies=technologies,
        verified=verified,
    )
    database.add(experience)
    database.commit()
    return experience


def add_project(
    database: Session,
    name: str,
    technologies: list[str],
) -> Project:
    project = Project(
        candidate_profile_id=1,
        name=name,
        description="A verified project.",
        technologies=technologies,
        features=[],
        achievements=[],
        verified=True,
    )
    database.add(project)
    database.commit()
    return project


def make_master_version(database: Session) -> CVVersion:
    now = datetime.now(timezone.utc)
    version = CVVersion(
        candidate_profile_id=1,
        version_type="verified_master",
        status="ready",
        display_name=f"Master {uuid4().hex[:8]}",
        is_primary=False,
        created_at=now,
        updated_at=now,
    )
    database.add(version)
    database.commit()
    database.refresh(version)
    return version


def make_job(
    database: Session,
    *,
    skills: list[str] | None = None,
    **overrides: object,
) -> NormalizedJob:
    source = create_job_source(
        JobSourceCreate(
            source_key=f"source-{uuid4().hex[:12]}",
            display_name="Example Robotics",
            source_category="provider_api",
        ),
        database,
    )
    payload: dict[str, object] = {
        "job_source_id": source.id,
        "source_job_id": f"job-{uuid4().hex[:16]}",
        "source_url": f"https://jobs.example.com/{uuid4().hex}",
        "employer_name": "Example Robotics",
        "title": "Platform Engineer",
        "content_hash": content_hash(),
        "normalization_version": "phase6-test",
    }
    payload.update(overrides)
    return create_normalized_job(
        NormalizedJobCreate(**payload),
        database,
        skill_requirements=[
            JobSkillRequirementCreate(
                skill_text=name,
                requirement_type="required",
            )
            for name in (skills or [])
        ],
    )


def request(key: str | None = None) -> CVMasterGenerationRequest:
    return CVMasterGenerationRequest(
        idempotency_key=key or f"phase6-{uuid4().hex[:16]}",
    )


def tailor(
    database: Session,
    job: NormalizedJob,
    key: str | None = None,
):
    return generate_tailored_cv(job.id, request(key), database)


def seed_vault(database: Session) -> None:
    make_profile(database)
    add_skill(database, "Python")
    add_skill(database, "Go")
    add_experience(
        database,
        "Python Corp",
        ["Python"],
        date(2020, 1, 1),
    )
    add_experience(
        database,
        "Legacy Corp",
        ["COBOL"],
        date(2015, 1, 1),
    )
    add_project(database, "API Platform", ["Python"])
    make_master_version(database)


def test_keyword_extraction_is_deterministic(
    db_session: Session,
) -> None:
    make_profile(db_session)
    python = add_skill(db_session, "Python")
    rust = add_skill(db_session, "Rust")
    job = make_job(
        db_session,
        skills=["Kubernetes"],
        description_text="We build services in Python. No rust here.",
    )

    job_skills = list(
        db_session.scalars(
            select(JobSkillRequirement).where(
                JobSkillRequirement.job_id == job.id
            )
        ).all()
    )
    keywords = extract_job_keywords(job, job_skills, [python, rust])

    assert "kubernetes" in keywords
    assert "python" in keywords
    assert "rust" in keywords


def test_experience_selection_prefers_keyword_overlap(
    db_session: Session,
) -> None:
    make_profile(db_session)
    relevant = add_experience(
        db_session,
        "Python Corp",
        ["Python"],
        date(2018, 1, 1),
    )
    add_experience(
        db_session,
        "Legacy Corp",
        ["COBOL"],
        date(2022, 1, 1),
    )
    experiences = list(db_session.scalars(select(Experience)).all())

    selected, _ = select_relevant_experiences(
        experiences,
        {"python"},
    )

    assert selected[0].id == relevant.id
    assert len(selected) == 2


def test_skill_reordering_puts_matches_first(
    db_session: Session,
) -> None:
    make_profile(db_session)
    add_skill(db_session, "Go")
    add_skill(db_session, "Python")
    skills = list(db_session.scalars(select(Skill)).all())

    ordered = reorder_skills(skills, {"python"})

    assert ordered[0].name == "Python"
    assert {skill.name for skill in ordered} == {"Python", "Go"}


def test_tailoring_requires_master_version(
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.setenv(
        "CAREERPILOT_GENERATED_CV_STORAGE_DIR",
        str(tmp_path / "generated"),
    )
    make_profile(db_session)
    job = make_job(db_session)

    with pytest.raises(CVTailoringError) as error:
        tailor(db_session, job)

    assert "master CV" in error.value.safe_message


def test_tailoring_generates_artifacts_and_link(
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.setenv(
        "CAREERPILOT_GENERATED_CV_STORAGE_DIR",
        str(tmp_path / "generated"),
    )
    seed_vault(db_session)
    job = make_job(db_session, skills=["Python"])

    result = tailor(db_session, job)

    assert result.created is True
    assert result.version.version_type == "tailored"
    assert result.version.status == "ready"
    assert result.version.job_reference is not None
    assert f"#{job.id}" in result.version.job_reference
    assert sorted(
        artifact.artifact_type for artifact in result.artifacts
    ) == ["docx", "pdf"]

    link = db_session.scalar(select(JobTailoredCV))
    assert link is not None
    assert link.job_id == job.id
    assert link.cv_version_id == result.version.id


def test_tailoring_is_idempotent_per_key(
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.setenv(
        "CAREERPILOT_GENERATED_CV_STORAGE_DIR",
        str(tmp_path / "generated"),
    )
    seed_vault(db_session)
    job = make_job(db_session, skills=["Python"])

    first = tailor(db_session, job, key="phase6-repeat-key")
    second = tailor(db_session, job, key="phase6-repeat-key")

    assert first.created is True
    assert second.created is False
    assert second.version.id == first.version.id


def test_same_key_for_different_jobs_creates_separate_cvs(
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.setenv(
        "CAREERPILOT_GENERATED_CV_STORAGE_DIR",
        str(tmp_path / "generated"),
    )
    seed_vault(db_session)
    job_a = make_job(db_session, skills=["Python"])
    job_b = make_job(db_session, skills=["Go"])

    first = tailor(db_session, job_a, key="phase6-shared-key")
    second = tailor(db_session, job_b, key="phase6-shared-key")

    assert first.version.id != second.version.id


def test_tailored_cv_contains_only_verified_facts(
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.setenv(
        "CAREERPILOT_GENERATED_CV_STORAGE_DIR",
        str(tmp_path / "generated"),
    )
    seed_vault(db_session)
    add_skill(db_session, "Kubernetes", verified=False)
    add_experience(
        db_session,
        "Unverified Corp",
        ["Kubernetes"],
        date(2023, 1, 1),
        verified=False,
    )
    job = make_job(db_session, skills=["Python", "Kubernetes"])

    result = tailor(db_session, job)

    docx_artifact = next(
        artifact
        for artifact in result.artifacts
        if artifact.artifact_type == "docx"
    )
    from docx import Document as DocxDocument

    from app.services.generated_cv_storage import (
        resolve_generated_artifact_path,
    )

    docx_path = resolve_generated_artifact_path(
        docx_artifact.storage_path,
    )
    text = "\n".join(
        paragraph.text
        for paragraph in DocxDocument(str(docx_path)).paragraphs
    )

    assert "Python Corp" in text
    assert "Unverified Corp" not in text
    assert "Kubernetes" not in text


def test_tailored_summary_adds_only_verified_skill_names(
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.setenv(
        "CAREERPILOT_GENERATED_CV_STORAGE_DIR",
        str(tmp_path / "generated"),
    )
    seed_vault(db_session)
    job = make_job(db_session, skills=["Python"])

    result = tailor(db_session, job)

    from docx import Document as DocxDocument

    from app.services.generated_cv_storage import (
        resolve_generated_artifact_path,
    )

    docx_artifact = next(
        artifact
        for artifact in result.artifacts
        if artifact.artifact_type == "docx"
    )
    text = "\n".join(
        paragraph.text
        for paragraph in DocxDocument(
            str(
                resolve_generated_artifact_path(
                    docx_artifact.storage_path,
                )
            )
        ).paragraphs
    )

    assert "Backend engineer focused on reliable APIs." in text
    assert "Verified skills relevant to this role: Python." in text


def test_tailoring_does_not_modify_vault_records(
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.setenv(
        "CAREERPILOT_GENERATED_CV_STORAGE_DIR",
        str(tmp_path / "generated"),
    )
    seed_vault(db_session)
    job = make_job(db_session, skills=["Python"])
    profile = db_session.get(CandidateProfile, 1)
    assert profile is not None
    summary_before = profile.professional_summary

    tailor(db_session, job)

    db_session.refresh(profile)
    assert profile.professional_summary == summary_before
    experiences = list(db_session.scalars(select(Experience)).all())
    assert {record.company for record in experiences} == {
        "Python Corp",
        "Legacy Corp",
    }


def test_retailoring_updates_job_link(
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.setenv(
        "CAREERPILOT_GENERATED_CV_STORAGE_DIR",
        str(tmp_path / "generated"),
    )
    seed_vault(db_session)
    job = make_job(db_session, skills=["Python"])

    first = tailor(db_session, job)
    second = tailor(db_session, job)

    assert second.version.id != first.version.id
    links = list(db_session.scalars(select(JobTailoredCV)).all())
    assert len(links) == 1
    assert links[0].cv_version_id == second.version.id


def test_get_tailored_cv_or_404(db_session: Session) -> None:
    with pytest.raises(CVTailoringError) as error:
        get_tailored_cv_or_404(987654, db_session)

    assert error.value.status_code == 404


def test_tailor_endpoints(
    client: TestClient,
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    monkeypatch.setenv(
        "CAREERPILOT_GENERATED_CV_STORAGE_DIR",
        str(tmp_path / "generated"),
    )
    seed_vault(db_session)
    job = make_job(db_session, skills=["Python"])

    missing = client.get(f"/api/jobs/{job.id}/tailored-cv")
    assert missing.status_code == 404

    created = client.post(
        f"/api/jobs/{job.id}/tailor-cv",
        json={"idempotency_key": "phase6-endpoint-key"},
    )
    assert created.status_code == 201
    body = created.json()
    assert body["created"] is True
    assert body["version"]["version_type"] == "tailored"
    assert len(body["artifacts"]) == 2
    for artifact in body["artifacts"]:
        assert "storage_path" not in artifact

    fetched = client.get(f"/api/jobs/{job.id}/tailored-cv")
    assert fetched.status_code == 200
    assert fetched.json()["version"]["id"] == body["version"]["id"]

    unknown = client.post(
        "/api/jobs/987654/tailor-cv",
        json={"idempotency_key": "phase6-unknown-job"},
    )
    assert unknown.status_code == 404


def test_phase6_migration_upgrade_downgrade_reupgrade(
    tmp_path: Path,
) -> None:
    backend_dir = Path(__file__).resolve().parents[1]
    migration_db = tmp_path / "phase6-migration.db"

    config = Config(str(backend_dir / "alembic.ini"))
    config.set_main_option("script_location", str(backend_dir / "alembic"))
    config.set_main_option("sqlalchemy.url", f"sqlite:///{migration_db}")

    command.upgrade(config, PHASE6_HEAD)
    engine = create_engine(f"sqlite:///{migration_db}")
    assert "job_tailored_cvs" in set(inspect(engine).get_table_names())
    engine.dispose()

    command.downgrade(config, PHASE5_HEAD)
    engine = create_engine(f"sqlite:///{migration_db}")
    tables = set(inspect(engine).get_table_names())
    assert "job_tailored_cvs" not in tables
    assert "job_matches" in tables
    engine.dispose()

    command.upgrade(config, PHASE6_HEAD)
    engine = create_engine(f"sqlite:///{migration_db}")
    assert "job_tailored_cvs" in set(inspect(engine).get_table_names())
    engine.dispose()
